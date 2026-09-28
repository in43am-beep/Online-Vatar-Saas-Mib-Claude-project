"""mib/gdrive.py -- Google Drive connector for Avatar Production by MIB.
"""
import re, threading
from datetime import datetime
from pathlib import Path

SCOPES = ["https://www.googleapis.com/auth/drive.file"]
_CONFIG_DIR = None
_CREDS_FILE = "gdrive_credentials.json"
_TOKEN_FILE  = "gdrive_token.json"
_SERVICE_ACCOUNT_FILE = "gdrive_service_account.json"  # server/headless path


class DriveError(Exception):
    pass


def init(config_dir):
    global _CONFIG_DIR
    _CONFIG_DIR = Path(config_dir)


def service_account_path():
    """Path of the service-account JSON, if present."""
    if not _CONFIG_DIR:
        return None
    p = _CONFIG_DIR / _SERVICE_ACCOUNT_FILE
    return p if p.is_file() else None


def is_configured():
    if not _CONFIG_DIR: return False
    return (service_account_path() is not None
            or ((_CONFIG_DIR / _TOKEN_FILE).is_file()
                and (_CONFIG_DIR / _CREDS_FILE).is_file()))


def _load_creds():
    try:
        # Server/headless path: service-account JSON (no browser needed).
        sa = service_account_path()
        if sa is not None:
            from google.oauth2 import service_account
            return service_account.Credentials.from_service_account_file(
                str(sa), scopes=SCOPES)
        from google.oauth2.credentials import Credentials
        from google.auth.transport.requests import Request
        tp = _CONFIG_DIR / _TOKEN_FILE
        if not tp.is_file(): raise DriveError("Not authenticated.")
        creds = Credentials.from_authorized_user_file(str(tp), SCOPES)
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
            tp.write_text(creds.to_json(), encoding="utf-8")
        return creds
    except DriveError: raise
    except Exception as e: raise DriveError(f"Drive error: {e}") from e


def connect_drive(creds_json_path=None, logger=None):
    try:
        from google_auth_oauthlib.flow import InstalledAppFlow
        cp = Path(creds_json_path) if creds_json_path else (_CONFIG_DIR / _CREDS_FILE)
        if not cp.is_file(): raise DriveError("No credentials file. Download from Google Cloud Console.")
        flow = InstalledAppFlow.from_client_secrets_file(str(cp), SCOPES)
        creds = flow.run_local_server(port=0)
        (_CONFIG_DIR / _TOKEN_FILE).write_text(creds.to_json(), encoding="utf-8")
        if logger: logger("Google Drive connected.")
        try:
            import googleapiclient.discovery as d
            svc = d.build("oauth2", "v2", credentials=creds)
            return svc.userinfo().get().execute().get("email", "connected")
        except Exception: return "connected"
    except DriveError: raise
    except Exception as e: raise DriveError(f"OAuth2 failed: {e}") from e


def _build_service():
    import googleapiclient.discovery as d
    return d.build("drive", "v3", credentials=_load_creds())


def _slug(text, n=50):
    s = re.sub(r"[^\w\s-]", "", str(text or "video"))
    return re.sub(r"[\s_-]+", "-", s).strip("-")[:n] or "video"


def _get_folder(svc, name, parent=None):
    n = name.replace(chr(39), "")
    q = f"name={chr(39)}{n}{chr(39)} and mimeType={chr(39)}application/vnd.google-apps.folder{chr(39)} and trashed=false"
    if parent: q += f" and {chr(39)}{parent}{chr(39)} in parents"
    hits = svc.files().list(q=q, fields="files(id)").execute().get("files", [])
    if hits: return hits[0]["id"]
    m = {"name": name, "mimeType": "application/vnd.google-apps.folder"}
    if parent: m["parents"] = [parent]
    return svc.files().create(body=m, fields="id").execute()["id"]


def _up(svc, lpath, dname, parent):
    from googleapiclient.http import MediaFileUpload
    lp = Path(lpath)
    if not lp.is_file(): return None
    ext = lp.suffix.lower()
    mime = {".mp4": "video/mp4", ".txt": "text/plain",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".json": "application/json"}.get(ext, "application/octet-stream")
    f = svc.files().create(body={"name": dname, "parents": [parent]},
        media_body=MediaFileUpload(str(lp), mimetype=mime, resumable=True), fields="id").execute()
    return f.get("id")


def _make_docx(job_dir, out_path):
    try:
        from docx import Document
        from docx.shared import Pt
        doc = Document()
        doc.add_heading("Video Metadata", 0)
        for h, fn in [("Title", "title.txt"), ("Description", "description.txt"),
                      ("Tags", "tags.txt"), ("Thumbnail Prompt", "thumbnail-prompt.txt")]:
            p = Path(job_dir) / fn
            if p.is_file():
                doc.add_heading(h, level=1)
                r = doc.add_paragraph(p.read_text(encoding="utf-8", errors="replace").strip()).runs
                if r: r[0].font.size = Pt(11)
        doc.save(str(out_path))
        return True
    except Exception: return False


def upload_job(job_dir, channel_name, title, final_mp4=None, logger=None):
    """Upload a finished video job folder to Google Drive. Never raises."""
    if not is_configured(): return {}
    try:
        svc  = _build_service()
        date = datetime.now().strftime("%Y-%m-%d")
        ch   = _get_folder(svc, str(channel_name or "MIB"))
        dt   = _get_folder(svc, date, ch)
        vf   = _get_folder(svc, _slug(title), dt)
        job  = Path(job_dir)
        done = {}
        if final_mp4 and Path(final_mp4).is_file():
            fid = _up(svc, final_mp4, "video.mp4", vf)
            if fid:
                done["video"] = fid
                if logger: logger("    gdrive: video.mp4 uploaded")
        for fn in ("title.txt", "description.txt", "tags.txt", "thumbnail-prompt.txt"):
            p = job / fn
            if p.is_file():
                fid = _up(svc, p, fn, vf)
                if fid: done[fn] = fid
        dp = job / "metadata.docx"
        if _make_docx(job_dir, dp):
            fid = _up(svc, dp, "metadata.docx", vf)
            if fid: done["docx"] = fid
        url = "https://drive.google.com/drive/folders/" + vf
        if logger: logger("    gdrive: done -> " + url)
        return {"folder_url": url, "folder_id": vf, "files": done}
    except DriveError as e:
        if logger: logger("    gdrive WARNING: " + str(e))
        return {}
    except Exception as e:
        if logger: logger("    gdrive WARNING: " + str(e))
        return {}


def account_info():
    """Return dict with email and storage info, or {} if not connected."""
    try:
        creds = _load_creds()
        import googleapiclient.discovery as d
        svc   = d.build("drive", "v3", credentials=creds)
        osvc  = d.build("oauth2", "v2", credentials=creds)
        info  = osvc.userinfo().get().execute()
        about = svc.about().get(fields="storageQuota").execute()
        q     = about.get("storageQuota", {})
        used  = int(q.get("usage", 0))
        total = int(q.get("limit", 1))
        return {"email": info.get("email", ""), "name": info.get("name", ""),
                "used_gb": round(used/(1024**3), 2),
                "total_gb": round(total/(1024**3), 2),
                "percent": round(used/max(total,1)*100, 1)}
    except Exception: return {}

